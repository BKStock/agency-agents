"""総務省「固定電話番号の指定状況」表の解析。

* XLS（xlrd）/ XLSX（openpyxl）/ CSV / ZIP（中の各ファイル）に対応
* 見出し語（config/layouts.json）で列を特定する。必須列が見つからなければ停止する
  （未知の表形式を推測で読まない）
* セルは文字列として扱う。数値セルの市外局番で先頭 0 が落ちていれば補い、補正として記録する
* 結合セル由来の空欄（市外局番・番号区画）は上の行から引き継ぎ、補正として記録する
* 基準日は表の記載（「令和6年4月1日現在」など）から読む
* 同じページに並ぶ固定電話以外の表（携帯・IP電話・フリーダイヤル等）は、題名が
  `out_of_scope_titles` の語を含むときだけ対象外として飛ばす。題名が一覧に無ければ未知の表形式として止める
"""

from __future__ import annotations

import csv
import hashlib
import io
import json
import sqlite3
import zipfile
from dataclasses import dataclass, field
from datetime import date
from pathlib import PurePosixPath

from ..db.connection import utc_now
from ..normalize import CodeError, compact, nfkc, parse_japanese_date, phone_part
from ..acquire import archive

TABLE_EXTS = (".xls", ".xlsx", ".csv")


class UnknownLayout(ValueError):
    pass


class OutOfScope(Exception):
    """固定電話以外の番号の表（題名で判定）。未知の表形式とは別に扱い、取込も停止もしない。"""

    def __init__(self, name: str, keyword: str):
        super().__init__(f"{name}: out of scope (title contains {keyword!r})")
        self.keyword = keyword


@dataclass
class ParsedTable:
    member: str
    as_of_date: date | None
    rows: list[dict]
    skipped_rows: int
    sheets_used: list[str] = field(default_factory=list)


def _sheets(name: str, body: bytes) -> list[tuple[str, list[list]]]:
    ext = PurePosixPath(name.lower()).suffix
    if ext == ".xlsx":
        from openpyxl import load_workbook  # 任意依存
        wb = load_workbook(io.BytesIO(body), read_only=True, data_only=True)
        out = [(ws.title, [list(r) for r in ws.iter_rows(values_only=True)]) for ws in wb.worksheets]
        wb.close()
        return out
    if ext == ".xls":
        import xlrd  # 任意依存（xls 専用）
        book = xlrd.open_workbook(file_contents=body)
        return [(sh.name, [sh.row_values(i) for i in range(sh.nrows)]) for sh in book.sheets()]
    if ext == ".csv":
        for enc in ("utf-8-sig", "cp932"):
            try:
                text = body.decode(enc)
                break
            except UnicodeDecodeError:
                continue
        else:
            raise UnknownLayout(f"{name}: cannot decode CSV")
        return [("csv", [row for row in csv.reader(io.StringIO(text))])]
    raise UnknownLayout(f"{name}: unsupported file type")


def _find_header(rows: list[list], layout: dict) -> tuple[int, dict[str, int]] | None:
    for i, row in enumerate(rows[: layout.get("header_search_rows", 40)]):
        cells = [compact(c) for c in row]
        used: set[int] = set()
        mapping: dict[str, int] = {}
        for col in layout["columns"]:
            for alias in col["aliases"]:
                hit = next((j for j, c in enumerate(cells) if j not in used and c and alias in c), None)
                if hit is not None:
                    mapping[col["name"]] = hit
                    used.add(hit)
                    break
        if all(c["name"] in mapping for c in layout["columns"] if c.get("required")):
            return i, mapping
    return None


def _out_of_scope_keyword(sheets: list[tuple[str, list[list]]], layout: dict) -> str | None:
    """見出しより上（題名の行）に、固定電話以外を表す語があればその語を返す。"""
    keywords = [(k, compact(k)) for k in layout.get("out_of_scope_titles", [])]
    for _, rows in sheets:
        for row in rows[: layout.get("header_search_rows", 40)]:
            for cell in row:
                text = compact(cell)
                hit = next((k for k, ck in keywords if ck and ck in text), None)
                if hit:
                    return hit
    return None


def _find_date(rows: list[list], upto: int) -> date | None:
    found = None
    for row in rows[: max(upto, 1) + 1]:
        for cell in row:
            text = nfkc(cell)
            if not text:
                continue
            d = parse_japanese_date(text)
            if d and "現在" in text:
                return d
            found = found or d
    return found


def parse_sheet(rows: list[list], layout: dict) -> tuple[date | None, list[dict], int] | None:
    found = _find_header(rows, layout)
    if found is None:
        return None
    header_row, mapping = found
    fill = {c["name"] for c in layout["columns"] if c.get("fill_down")}
    carry: dict[str, object] = {}
    out, skipped = [], 0
    for raw in rows[header_row + 1:]:
        cells = {name: (raw[j] if j < len(raw) else None) for name, j in mapping.items()}
        if all(nfkc(v) == "" for v in cells.values()):
            continue
        repaired = []
        for name in fill:
            if name in cells and nfkc(cells[name]) == "" and name in carry:
                cells[name] = carry[name]
                repaired.append(f"fill_down:{name}")
            elif name in cells and nfkc(cells[name]) != "":
                carry[name] = cells[name]
        try:
            area, fixed = phone_part(cells["area_code"], "area_code", leading_zero=True)
            local, _ = phone_part(cells["local_code"], "local_code", leading_zero=False)
        except CodeError:
            skipped += 1  # 注記行・小計行など
            continue
        if fixed:
            repaired.append("leading_zero:area_code")
        rec = {"area_code": area, "local_code": local, "repaired": ";".join(repaired) or None}
        for name in ("number_area_code", "number_area_name", "carrier", "usage_status"):
            v = cells.get(name)
            if isinstance(v, float) and v.is_integer():
                v = int(v)
            rec[name] = (str(v).strip() or None) if v is not None else None  # 原文のまま（比較時に正規化）
        out.append(rec)
    return _find_date(rows, header_row), out, skipped


def parse_file(name: str, body: bytes, layout: dict) -> list[ParsedTable]:
    """1つの原本から ParsedTable の列を返す（ZIP なら中の表ファイルごと）。"""
    if name.lower().endswith(".zip"):
        tables, out_of_scope = [], None
        with zipfile.ZipFile(io.BytesIO(body)) as zf:
            for info in sorted(zf.infolist(), key=lambda i: i.filename):
                if info.is_dir() or not info.filename.lower().endswith(TABLE_EXTS):
                    continue
                try:
                    tables += [_with_member(t, info.filename) for t in parse_file(info.filename, zf.read(info), layout)]
                except OutOfScope as exc:
                    out_of_scope = exc
        if not tables and out_of_scope:
            raise OutOfScope(name, out_of_scope.keyword)
        if not tables:
            raise UnknownLayout(f"{name}: no table files inside ZIP")
        return tables
    parsed_rows, dates, used, skipped = [], [], [], 0
    sheets = _sheets(name, body)
    for sheet_name, rows in sheets:
        result = parse_sheet(rows, layout)
        if result is None:
            continue
        d, recs, sk = result
        used.append(sheet_name)
        dates.append(d)
        parsed_rows += recs
        skipped += sk
    if not used:
        keyword = _out_of_scope_keyword(sheets, layout)
        if keyword:
            raise OutOfScope(name, keyword)
        raise UnknownLayout(f"{name}: required headers {[c['aliases'][0] for c in layout['columns'] if c.get('required')]}"
                            " not found; update config/layouts.json after checking the file")
    known = sorted({d for d in dates if d})
    if len(known) > 1:
        raise UnknownLayout(f"{name}: sheets carry different as-of dates {known}")
    return [ParsedTable("", known[0] if known else None, parsed_rows, skipped, used)]


def _with_member(t: ParsedTable, member: str) -> ParsedTable:
    t.member = member if not t.member else f"{member}/{t.member}"
    return t


def content_hash(rows: list[dict]) -> str:
    keys = ("area_code", "local_code", "number_area_code", "number_area_name", "carrier", "usage_status")
    norm = sorted(tuple(compact(r.get(k)) for k in keys) for r in rows)
    return hashlib.sha256(json.dumps(norm, ensure_ascii=False).encode()).hexdigest()


# 地域に属さない番号（フリーダイヤル・IP電話・携帯など）。地域の系統に混ぜず、別の系統にする
NON_GEOGRAPHIC = ("0120", "0800", "0570", "0180", "0990", "020", "050", "060", "070", "080", "090")


def region_series(area_code: str, non_geographic: tuple[str, ...] = NON_GEOGRAPHIC) -> str:
    """系統＝市外局番の先頭2桁（01〜09 の地域）。ファイル名の数字は使わない。

    地域に属さない番号は、その番号の頭（例 0120）をそのまま系統にする。
    """
    for prefix in sorted(non_geographic, key=len, reverse=True):
        if area_code.startswith(prefix):
            return prefix
    return area_code[:2]


def ingest(conn: sqlite3.Connection, raw_dir, sha: str, layout: dict) -> list[int]:
    """原本を解析して version と designation_row を追記する。解析済みなら既存の version_id を返す。

    1つの原本（ZIP なら中の1ファイル）に複数の地域が含まれていれば、地域ごとに別の版にする。
    こうしておけば、ファイルの分け方が年によって変わっても、同じ地域同士で比較できる。
    """
    existing = [r["version_id"] for r in conn.execute(
        "SELECT version_id FROM numbers.version WHERE sha256 = ? ORDER BY version_id", (sha,))]
    if existing:
        return existing
    src = conn.execute("SELECT filename, wayback_ts FROM numbers.raw_source WHERE sha256 = ?"
                       " ORDER BY origin = 'official' DESC, fetched_at LIMIT 1", (sha,)).fetchone()
    body = archive.read(conn, raw_dir, sha)
    ids = []
    for table in parse_file(src["filename"] or "file", body, layout):
        if table.as_of_date:
            as_of, date_source = table.as_of_date.isoformat(), "content"
        elif src["wayback_ts"]:
            ts = src["wayback_ts"]
            as_of, date_source = f"{ts[:4]}-{ts[4:6]}-{ts[6:8]}", "wayback_ts"
        else:
            as_of, date_source = None, "unknown"
        groups: dict[str, list[dict]] = {}
        nongeo = tuple(layout.get("non_geographic_prefixes", NON_GEOGRAPHIC))
        for r in table.rows:
            groups.setdefault(region_series(r["area_code"], nongeo), []).append(r)
        for series, rows in sorted(groups.items()):
            vid = conn.execute(
                "INSERT INTO numbers.version (sha256, member, series, as_of_date, date_source, layout, row_count,"
                " content_sha256, parsed_at) VALUES (?,?,?,?,?,?,?,?,?)",
                (sha, table.member, series, as_of, date_source, layout["layout_id"], len(rows),
                 content_hash(rows), utc_now())).lastrowid
            conn.executemany(
                "INSERT INTO numbers.designation_row VALUES (?,?,?,?,?,?,?,?,?)",
                [(vid, i, r["area_code"], r["local_code"], r["number_area_code"], r["number_area_name"],
                  r["carrier"], r["usage_status"], r["repaired"]) for i, r in enumerate(rows)])
            ids.append(int(vid))
    return ids
