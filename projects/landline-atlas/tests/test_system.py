"""統合：パイプライン・検索・出力・画面・調査優先順位・保護パス・バックアップ（受入 4, 10, 11, 13）。"""

import csv
import json
import tempfile
import threading
import unittest
import urllib.request
from pathlib import Path

from support import (BV, V2023, V2024, V2025, YEAR, FakeTransport, db, official_routes, rules, seed_regions,
                     wayback_routes, xlsx)

from landline_atlas import paths
from landline_atlas.classify import engine
from landline_atlas.db.connection import connect
from landline_atlas.evidence import facilities
from landline_atlas.export.writer import export_all
from landline_atlas.pipeline import update
from landline_atlas.qa import backup
from landline_atlas.search.query import Filter, search
from landline_atlas.web import server


def add_area_map(conn):
    conn.execute("INSERT INTO numbers.number_area_municipality VALUES ('A27-1', '27100', '06', '大阪', 'part',"
                 " '2020-01-01', NULL)")


class PipelineTest(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.root = Path(self.tmp.name) / "atlas"
        self.paths = paths.DataPaths(self.root).ensure()
        self.local = Path(self.tmp.name) / "local_2023.xlsx"
        self.local.write_bytes(xlsx("令和5年4月1日現在", V2023))
        before = self.local.read_bytes()
        body25, body24 = xlsx("令和7年1月1日現在", V2025), xlsx("令和6年4月1日現在", V2024)
        self.transport = FakeTransport({**official_routes(body25),
                                        **wayback_routes({"20240410000000": body24, "20250110000000": body25})})
        conn = connect(self.paths.db)
        seed_regions(conn)
        add_area_map(conn)
        facilities.add(conn, facility_name="市立テスト小学校", facility_type="school", area_code="06",
                       local_code="6200", address="架空市戸建町", town="27100000101", census_year=YEAR,
                       boundary_version=BV, verified=True, source_url="https://example.invalid/school",
                       observed_on="2026-10-01")
        conn.commit()
        conn.close()
        self.manifest = update(self.paths, wayback=True, local_files=[self.local], transport=self.transport,
                               sleep=lambda s: None)
        self.assertEqual(self.local.read_bytes(), before)  # 手元の原本は変更しない

    def test_manifest_and_history(self):
        m = self.manifest
        self.assertEqual(m["outcome"], "ok", json.dumps(m, ensure_ascii=False, indent=1))
        h = m["stages"]["history"]
        self.assertEqual((h["series"], h["selected_versions"], h["pairs"]), (1, 3, 2))
        self.assertEqual((h["change_rows_gross"], h["distinct_codes_became_in_use"]), (5, 2))
        self.assertEqual(m["stages"]["acquire"]["new_raw_files"], 3)
        self.assertTrue(m["backup"]["ok"])
        self.assertTrue(all(v["passed"] for v in m["stages"]["qa"].values() if v["severity"] == "error"))
        runs = sorted(self.paths.runs.iterdir())
        self.assertTrue((runs[-1] / "run_manifest.json").exists())
        export_dir = Path(m["stages"]["export"]["dir"])
        self.assertTrue((export_dir / "landline_atlas.xlsx").exists())
        with open(export_dir / "09b_change_summary.csv", encoding="utf-8-sig") as fh:
            summary = {r["metric"]: r["value"] for r in csv.DictReader(fh)}
        self.assertEqual((summary["change_rows_gross"], summary["distinct_codes_became_in_use"]), ("5", "2"))
        conn = connect(self.paths.db, read_only=True)
        sel = {r["as_of_date"]: r["reason"] for r in conn.execute("SELECT * FROM numbers.selected_version")}
        # Wayback の 2025 版は公式取得と同一バイト列 → 原本1つ・入手経路2つ・版1つ
        self.assertEqual(sel["2025-01-01"], "single_candidate")
        self.assertEqual(conn.execute(
            "SELECT COUNT(*) FROM numbers.raw_source s JOIN numbers.version v ON v.sha256 = s.sha256"
            " WHERE v.as_of_date = '2025-01-01'").fetchone()[0], 2)
        conn.close()

    def test_second_run_is_idempotent(self):
        m2 = update(self.paths, wayback=True, transport=self.transport, sleep=lambda s: None, do_export=False)
        self.assertEqual(m2["outcome"], "ok")
        self.assertEqual(m2["stages"]["acquire"]["new_raw_files"], 0)
        self.assertEqual((m2["stages"]["ingest"]["new_versions"], m2["stages"]["history"]["new_pairs"]), (0, 0))

    def test_unknown_layout_blocks_classification(self):
        from openpyxl import Workbook
        import io
        wb = Workbook()
        wb.active.append(["電話番号", "地域"])
        buf = io.BytesIO()
        wb.save(buf)
        bad = Path(self.tmp.name) / "bad.xlsx"
        bad.write_bytes(buf.getvalue())
        m = update(self.paths, online=False, local_files=[bad], do_export=False)
        self.assertEqual(m["outcome"], "blocked_by_qa")
        self.assertFalse(m["stages"]["qa"]["Q00_unknown_layout"]["passed"])
        self.assertNotIn("classify", m["stages"])

    def test_out_of_scope_table_does_not_block(self):
        from openpyxl import Workbook
        import io
        wb = Workbook()
        for row in (["電気通信番号指定状況"], ["・ＩＰ電話の番号（０５０－ＣＤＥＦ－ＧＨＪＫ）"], ["番号", "０", "１"]):
            wb.active.append(row)
        buf = io.BytesIO()
        wb.save(buf)
        ip = Path(self.tmp.name) / "050.xlsx"
        ip.write_bytes(buf.getvalue())
        m = update(self.paths, online=False, local_files=[ip], do_export=False)
        self.assertEqual(m["outcome"], "ok")
        self.assertEqual(m["stages"]["ingest"]["unknown_layouts"], [])
        self.assertEqual([f["keyword"] for f in m["stages"]["ingest"]["out_of_scope"]], ["IP電話"])
        self.assertTrue(m["stages"]["qa"]["Q00b_out_of_scope"]["passed"])
        self.assertNotIn("Q00_unknown_layout", m["stages"]["qa"])

    def test_search_relations_and_priority(self):  # 受入4, 11
        conn = connect(self.paths.db, read_only=True)
        hidden = search(conn, Filter(prefecture="27", mode="residential"))
        shown = search(conn, Filter(prefecture="27", mode="residential", include_r6=True))
        self.assertNotIn("R6", {r["class_code"] for r in hidden})
        self.assertIn("R6", {r["class_code"] for r in shown})
        both = search(conn, Filter(classes=["L1", "L6"], all_classes=True))
        self.assertEqual({r["town_key"] for r in both}, {"03215000100"})
        res = search(conn, Filter(town="27100000101", local_code="6200"))
        levels = {x["relation_level"] for x in res[0]["phone_relations"]}
        self.assertEqual(levels, {"area", "observed_verified"})
        res = search(conn, Filter(town="27100000301", local_code="6202", carrier="KDDI"))
        levels = {x["relation_level"] for x in res[0]["phone_relations"]}
        self.assertEqual(levels, {"area", "no_observation"})
        items = {(r["local_code"]): dict(r) for r in conn.execute(
            "SELECT * FROM priority_item WHERE priority_run_id = (SELECT MAX(priority_run_id) FROM priority_run)")}
        self.assertEqual(items["6200"]["evidence_level"], "observed_verified")
        self.assertEqual(items["6202"]["evidence_level"], "designation_only")
        self.assertGreater(items["6202"]["score"], items["6200"]["score"])
        self.assertEqual(set(items), {"6200", "6201", "6202"})  # 指定済み・使用中の局番だけ
        conn.close()

    def test_web_pages(self):
        conn = connect(self.paths.db, read_only=True)
        default = server.render_search(conn, {"pref": ["27"]})
        self.assertIn("R6（住宅・商業混在地区）を初期設定で非表示", default)
        self.assertNotIn("<td>R6", default)
        self.assertIn("<td>R6", server.render_search(conn, {"pref": ["27"], "r6": ["1"]}))
        self.assertIn("観測・確認済み", server.render_town(conn, "27100000101"))
        code = server.render_code(conn, "06", "6201")
        self.assertIn("carrier_changed", code)
        self.assertIn("ソフトバンク", code)
        self.assertIn("調査優先順位", server.render_priority(conn))
        conn.close()
        httpd = server.serve(self.paths.db, 0)
        t = threading.Thread(target=httpd.serve_forever, daemon=True)
        t.start()
        try:
            port = httpd.server_address[1]
            with urllib.request.urlopen(f"http://127.0.0.1:{port}/dashboard", timeout=10) as r:
                self.assertIn("都道府県別ダッシュボード", r.read().decode())
        finally:
            httpd.shutdown()
            httpd.server_close()


class ExportTest(unittest.TestCase):
    def test_reconciliation_and_counts(self):  # 受入10
        with tempfile.TemporaryDirectory() as tmp:
            conn = db(Path(tmp))
            seed_regions(conn)
            engine.classify(conn, rules(), YEAR, BV)
            res = export_all(conn, Path(tmp) / "out")
            rec = res["reconciliation"]
            self.assertEqual(rec["per_class"]["R2"], 2)
            self.assertGreater(rec["gross_class_matches"], rec["unique_towns_matched"])
            with open(Path(tmp) / "out" / "01_all_regions_master.csv", encoding="utf-8-sig") as fh:
                master = list(csv.DictReader(fh))
            self.assertEqual(master[0]["town_key"], "03215000100")  # 先頭ゼロ保持
            self.assertEqual([m for m in master if m["town_key"] == "27100000801"][0]["matched_classes"], "0")
            from openpyxl import load_workbook
            ws = load_workbook(Path(tmp) / "out" / "landline_atlas.xlsx")["01_all_regions_master"]
            self.assertEqual(ws["C2"].value, "03215000100")
            conn.close()


class SafetyTest(unittest.TestCase):
    def test_protected_legacy_directory(self):
        with self.assertRaises(paths.ProtectedPathError):
            paths.resolve("~/official_phone_data/new")
        with self.assertRaises(paths.ProtectedPathError):
            paths.check_writable_location(Path("/tmp/legacy/x"), protected=(Path("/tmp/legacy"),))

    def test_facility_rules(self):
        with tempfile.TemporaryDirectory() as tmp:
            conn = db(Path(tmp))
            with self.assertRaises(ValueError):
                facilities.add(conn, facility_name="個人宅", facility_type="residence", area_code="06",
                               local_code="6200", address="x", source_url="https://example.invalid",
                               observed_on="2026-10-01")
            with self.assertRaises(ValueError):
                facilities.add(conn, facility_name="x", facility_type="school", area_code="06", local_code="6200",
                               address="x", source_url="", observed_on="2026-10-01")
            conn.close()

    def test_backup_restore(self):  # 受入13
        with tempfile.TemporaryDirectory() as tmp:
            conn = db(Path(tmp))
            seed_regions(conn)
            conn.commit()
            conn.close()
            r = backup.backup(Path(tmp) / "db", Path(tmp) / "backups")
            self.assertTrue(r["ok"])
            self.assertEqual(set(r["databases"]), {"numbers", "regions", "evidence", "analysis"})


if __name__ == "__main__":
    unittest.main()
